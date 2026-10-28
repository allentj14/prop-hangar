"""Build the Excel workbook from planes2.json and lineups.json."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open(sys.argv[1]))
lu = json.load(open(sys.argv[2]))
OUT = sys.argv[3]
NATIONS = ["USA", "Germany", "USSR", "Great Britain", "Japan", "China", "Italy", "France", "Sweden", "Israel"]
ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V", 6: "VI", 7: "VII", 8: "VIII", 9: "IX"}
nb = lambda s: (s or "").replace("\xa0", " ")

for r in d:
    for k in ("name", "full_name", "guns", "def_guns", "bombs", "rockets", "torpedoes", "gun_pods", "folder"):
        r[k] = nb(r.get(k))
import re
RND = {"climb_rate_ms": 1, "climb_time_s": 0, "turn_time_s": 1, "roll_rate": 0, "ceiling_m": 0, "climb_alt_m": 0,
       "wing_loading": 0, "pwr": 2, "takeoff_mass_kg": 0, "empty_mass_kg": 0, "fuel_kg": 0, "vne_kmh": 0, "rp_mult": 2, "sl_mult_rb": 2}
for r in d:
    for k, p in RND.items():
        if isinstance(r.get(k), (int, float)): r[k] = round(r[k], p) if p else int(round(r[k]))
    for k in ("gun_ammo", "def_ammo", "repair_ab", "repair_rb", "cost_sl"):
        if not r.get(k): r[k] = None
    em = r.get("engine_model", "").replace("_", " ")
    parts = em.split(" ", 1)
    if len(parts) == 2 and parts[1].lower().startswith(parts[0].lower()): em = parts[1]
    r["engine_model"] = em.strip()
d.sort(key=lambda r: (NATIONS.index(r["nation"]), r["rank"], r["column"], r["order"]))

COLS = [
    ("Nation", "nation", 12), ("Rank", lambda r: ROMAN[r["rank"]], 6), ("Tree column", "column", 7),
    ("Tree position", "order", 7), ("Folder", "folder", 14), ("Plane", "name", 22), ("Full name", "full_name", 26),
    ("How to get", "obtain", 16), ("Role", "role", 16), ("Play style (community)", "style", 26), ("Play style (from stats)", "stat_style", 22), ("Tags (bombers & attackers)", lambda r: "; ".join(r.get("tags") or []), 40), ("Max bomb load (kg)", "bomb_load_kg", 8),
    ("BR AB", "br_ab", 6), ("BR RB", "br_rb", 6), ("BR SB", "br_sb", 6),
    ("Top speed (km/h)", "top_speed_kmh", 9), ("At altitude (m)", "top_speed_alt_m", 9),
    ("Sea-level speed (km/h)", "sl_speed_kmh", 9), ("Climb rate (m/s)", "climb_rate_ms", 8),
    ("Time to climb (s)", "climb_time_s", 8), ("Climb to (m)", "climb_alt_m", 8),
    ("Turn time (s)", "turn_time_s", 7), ("Roll rate (°/s)", "roll_rate", 7), ("Ceiling (m)", "ceiling_m", 8),
    ("Wing loading (kg/m²)", "wing_loading", 8), ("Power/weight", "pwr", 7),
    ("Dive limit Vne (km/h)", "vne_kmh", 8), ("Flap rip speeds (km/h)", "flap_limits_kmh", 12),
    ("Forward guns", "guns", 40), ("Forward ammo", "gun_ammo", 8), ("Burst mass (kg/s)", "burst_kg_s", 8),
    ("Defensive guns", "def_guns", 34), ("Defensive ammo", "def_ammo", 8),
    ("Bombs (options)", "bombs", 40), ("Rockets (options)", "rockets", 30), ("Torpedoes", "torpedoes", 24),
    ("Optional / alternate guns", "gun_pods", 26),
    ("Crew", "crew", 5), ("Engines", "engines", 6), ("Engine", "engine_model", 18), ("Engine type", "engine_type", 9),
    ("Takeoff mass (kg)", "takeoff_mass_kg", 8), ("Empty mass (kg)", "empty_mass_kg", 8), ("Fuel (kg)", "fuel_kg", 7),
    ("Max flight time (min)", "flight_time_min", 7), ("Takeoff run (m)", "takeoff_run_m", 7),
    ("Research (RP)", "research_rp", 9), ("Cost (SL)", "cost_sl", 9), ("Cost (GE)", "cost_ge", 8),
    ("Repair AB (SL)", "repair_ab", 8), ("Repair RB (SL)", "repair_rb", 8), ("RP multiplier", "rp_mult", 7),
    ("SL multiplier RB", "sl_mult_rb", 7),
    ("Strengths (vs same-BR planes)", "strengths", 40), ("Weaknesses (vs same-BR planes)", "weaknesses", 40),
    ("Community: pros", "comm_pros", 40), ("Community: cons", "comm_cons", 40), ("Community verdict", "comm_verdict", 26),
    ("How to fly it", "tips", 60), ("Community source", "comm_src", 40),
    ("Verification", "verify_status", 16), ("Verification notes", "verify_note", 50), ("Verification sources", "verify_src", 50),
]
# show armament right after play style so it's visible without scrolling
_front = [c for c in COLS if c[1] in ("guns", "gun_ammo", "def_guns", "bombs", "rockets")]
_rest = [c for c in COLS if c not in _front]
_i = [c[1] for c in _rest].index("stat_style") + 1
COLS = _rest[:_i] + _front + _rest[_i:]
NUMFMT = {"cost_sl": "#,##0", "research_rp": "#,##0", "cost_ge": "#,##0", "repair_ab": "#,##0", "repair_rb": "#,##0",
          "takeoff_mass_kg": "#,##0", "empty_mass_kg": "#,##0", "ceiling_m": "#,##0", "top_speed_alt_m": "#,##0",
          "br_ab": "0.0", "br_rb": "0.0", "br_sb": "0.0", "gun_ammo": "#,##0", "def_ammo": "#,##0"}

F = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F2A36")
HDR_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
STYLE_FILL = {"Boom & Zoom": "F6D7C3", "Energy fighter": "D6E8F7", "Turn fighter": "D7EFD5", "Attacker": "EFE3C8"}
BOMB_FILL = "E5E1EE"
RANK_FILL = PatternFill("solid", fgColor="EEF1F4")
thin = Side(style="thin", color="D0D5DB")

def style_fill(s):
    for k, v in STYLE_FILL.items():
        if s.startswith(k): return PatternFill("solid", fgColor=v)
    if "bomber" in s.lower(): return PatternFill("solid", fgColor=BOMB_FILL)
    return None

def val(r, k):
    v = k(r) if callable(k) else r.get(k)
    if v in ("", None): return None
    if isinstance(v, float) and v.is_integer() and k not in ("br_ab", "br_rb", "br_sb"): v = int(v)
    return v

def plane_sheet(ws, rows):
    ws.append([c[0] for c in COLS])
    for i, c in enumerate(COLS, 1):
        cell = ws.cell(1, i); cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = c[2]
    ws.row_dimensions[1].height = 45
    si = [c[1] for c in COLS].index("style") + 1
    for r in rows:
        ws.append([val(r, c[1]) for c in COLS])
        rr = ws.max_row
        for i, c in enumerate(COLS, 1):
            cell = ws.cell(rr, i)
            cell.font = BOLD if c[1] == "name" else BODY
            cell.border = Border(bottom=thin)
            if isinstance(c[1], str) and c[1] in NUMFMT: cell.number_format = NUMFMT[c[1]]
            if c[2] >= 24: cell.alignment = Alignment(wrap_text=True, vertical="top")
            else: cell.alignment = Alignment(vertical="top")
        f = style_fill(r["style"])
        if f: ws.cell(rr, si).fill = f
        if r["rank"] % 2 == 0:
            for i in (2, 3, 4): ws.cell(rr, i).fill = RANK_FILL
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"

wb = Workbook()
ws = wb.active; ws.title = "Read me"
info = [
    ("War Thunder propeller aircraft — all nations", True),
    ("", False),
    (f"{len(d)} propeller aircraft across 10 nations, including premium, event and squadron planes. Jets, rocket planes, turboprops and helicopters are left out.", False),
    ("Data source: War Thunder game files (community datamine of game version 2.59.0.90, 27 Oct 2026): research tree layout, battle ratings, flight-model and weapon files. Battle ratings change with each BR update.", False),
    ("", False),
    ("How the sheets are laid out", True),
    ("• One tab per nation, plus 'All planes'. Rows follow the research tree: rank (I, II, III…), then tree column left to right, then position top to bottom in that column. Folder = the group the plane sits in inside the tree.", False),
    ("• 'Best lineups' gives a recommended 5-plane lineup for every battle rating each nation can field, for Realistic (RB) and Arcade (AB). A tech-tree-only version is next to it in case you don't own the premiums.", False),
    ("• Every sheet has filters on the header row — use them to show only one rank, style or BR.", False),
    ("", False),
    ("What the stats mean", True),
    ("• Top speed, climb, turn time, roll rate and ceiling are the reference values shown on the in-game stat card (upgraded aircraft). Sea-level speed comes from the game's performance tables. Actual numbers in a match vary with mode, fuel and modifications.", False),
    ("• Turn time = seconds for a full 360° sustained turn (lower is better). Burst mass = kg of ammunition fired per second by forward guns — a quick way to compare firepower.", False),
    ("• Dive limit (Vne) = speed at which the wings tear off. Flap rip speeds = combat / takeoff / landing flap limits.", False),
    ("• Bomb, rocket and torpedo columns list the options available. 'up to N×' is the most of that weapon the plane can carry in one loadout; different options usually can't all be carried at once.", False),
    ("", False),
    ("Play style", True),
    ("• Boom & Zoom (BnZ): fast, heavy, dives well, turns poorly. Dive from above, shoot, climb away.", False),
    ("• Energy fighter: strong climb and engine power with decent handling. Keep speed and height, fight in the vertical.", False),
    ("• Turn fighter: light, low wing loading, tight turns. Wins slow, close dogfights, loses long chases.", False),
    ("• Two styles (e.g. 'Energy fighter / Turn fighter') means it's almost equally good at both. Bombers and attackers are labeled by role; if their guns and speed are good enough to fight, the style is added.", False),
    ("• 'Play style (community)' is what the player-written War Thunder wiki articles recommend for that plane or its close variants (562 of the 601 fighters and attackers). Where no community article exists, and for bombers, it falls back to the stats-based label.", False),
    ("• 'Play style (from stats)' compares the plane's speed, climb, turn, wing loading, dive limit and firepower against every prop within ±0.7 BR. When the two disagree, the community label is usually more reliable; the stats label shows how the plane compares on paper.", False),
    ("• Community pros, cons, verdict and 'How to fly it' tips are summarized from the wiki article linked in 'Community source'. Variants of one family (e.g. all P-47Ds) share one article's notes. Strengths/weaknesses columns are the stats comparison.", False),
    ("", False),
    ("Verification", True),
    ("• The 'Verification' column shows how each plane's label was checked: Confirmed (a second community source agrees), Mixed opinions (sources disagree; notes explain), Adjusted (label changed to match the community), Wiki + stats agree, Wiki only, or Game data only (no community writeup found, mostly niche bombers and event planes). Sources: War Thunder wiki, official forum threads, Steam community guides, and YouTube review titles.", False),
    ("", False),
    ("Lineups", True),
    ("• Picked by overall performance against the planes you'll meet at that BR (speed, climb, turn, firepower, roll), nudged up or down when the community calls a plane strong or weak for its BR, with a penalty for planes far below the lineup BR. Arcade lineups include one strike plane for ground targets; Realistic lineups are the five strongest fighters, ranked — in Air RB you usually get one life, so slot 1 is your main pick.", False),
]
for t, b in info:
    ws.append([t]); c = ws.cell(ws.max_row, 1)
    c.font = Font(name=F, bold=b, size=14 if ws.max_row == 1 else (11 if b else 10)); c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 120
for k, v in STYLE_FILL.items():
    ws.append([k]); c = ws.cell(ws.max_row, 1); c.fill = PatternFill("solid", fgColor=v); c.font = BODY
ws.append(["Bombers"]); c = ws.cell(ws.max_row, 1); c.fill = PatternFill("solid", fgColor=BOMB_FILL); c.font = BODY

plane_sheet(wb.create_sheet("All planes"), d)
for n in NATIONS:
    plane_sheet(wb.create_sheet(n), [r for r in d if r["nation"] == n])

ws = wb.create_sheet("Best lineups", 1)
hdr = ["Nation", "Mode", "Lineup BR"] + [f"Pick {i}" for i in range(1, 6)] + [f"Tech tree only {i}" for i in range(1, 6)]
ws.append(hdr)
for i, h in enumerate(hdr, 1):
    c = ws.cell(1, i); c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(i)].width = 12 if i <= 3 else 26
lu.sort(key=lambda x: (NATIONS.index(x["nation"]), 0 if x["mode"] == "RB" else 1, x["br"]))
PREM = PatternFill("solid", fgColor="FBEFC9")
for x in lu:
    fmt = lambda p: f"{nb(p['name'])} ({p['br']:.1f}, {p['style']})" + ("" if p["obtain"] == "Tech tree" else " ★")
    row = [x["nation"], x["mode"], x["br"]] + [fmt(p) for p in x["best"]] + [""] * (5 - len(x["best"])) \
          + ([fmt(p) for p in x["tech"]] + [""] * (5 - len(x["tech"])) if not x["same"] else ["Same as picks"] + [""] * 4)
    ws.append(row)
    rr = ws.max_row
    for i in range(1, len(hdr) + 1):
        c = ws.cell(rr, i); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(bottom=thin)
    ws.cell(rr, 3).number_format = "0.0"
    for j, p in enumerate(x["best"]):
        if p["obtain"] != "Tech tree": ws.cell(rr, 4 + j).fill = PREM
ws.append([]); ws.append(["★ = premium, event or squadron plane (highlighted). Style: BnZ = Boom & Zoom, Energy = energy fighter, Turn = turn fighter, Atk = attacker."])
ws.cell(ws.max_row, 1).font = Font(name=F, size=10, italic=True)
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{len(lu) + 1}"
ws = wb.create_sheet("Lineups by play style", 2)
hdr = ["Nation", "Mode", "Lineup BR", "Play style"] + [f"Pick {i}" for i in range(1, 6)]
ws.append(hdr)
for i, h in enumerate(hdr, 1):
    c = ws.cell(1, i); c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(i)].width = 12 if i <= 4 else 26
SN = {"bnz": "Boom & Zoom", "nrg": "Energy fighter", "trn": "Turn fighter"}
for x in lu:
    for k, name in SN.items():
        picks = x.get("styles", {}).get(k) or []
        if not picks: continue
        row = [x["nation"], x["mode"], x["br"], name] + [fmt(p) for p in picks] + [""] * (5 - len(picks))
        ws.append(row)
        rr = ws.max_row
        for i in range(1, len(hdr) + 1):
            c = ws.cell(rr, i); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(bottom=thin)
        ws.cell(rr, 3).number_format = "0.0"
        f = style_fill(name)
        if f: ws.cell(rr, 4).fill = f
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{ws.max_row}"
wb.save(OUT)
print("saved", OUT)
